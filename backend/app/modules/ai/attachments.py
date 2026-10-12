"""Turn an attachment into units the local model can read within its limits (16k context, 2k output).

- PDF: one unit per page; the text layer when the page has one, otherwise the page rendered as an image (scans).
- Images and screenshots: downscaled; tall phone screenshots are cut into overlapping tiles.
- CSV, TSV, TXT and Excel: text chunks of whole lines, each repeating the header line.
A unit can be split again (halves) when the model runs out of output room.
"""

import base64
import csv
import io
import re
from collections.abc import Callable
from dataclasses import dataclass

MAX_BYTES = 25 * 1024 * 1024
MAX_UNITS = 60
CHUNK_CHARS = 6000
LONG_SIDE = 1600
TILE_RATIO = 1.4  # tile height relative to width for tall screenshots
MIN_TEXT = 15  # fewer letters/digits than this on a PDF page → treat it as a scan (short pages still count)

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
TEXT_EXT = {".csv", ".tsv", ".txt"}
SHEET_EXT = {".xlsx", ".xlsm"}


class AttachmentError(ValueError):
    pass


@dataclass
class Unit:
    label: str  # "page 3", "image tile 2/4", "text lines 1-120"
    text: str | None = None
    image: bytes | None = None  # JPEG
    as_image: Callable[[], list["Unit"]] | None = None  # PDF text page: render it when its text is unusable

    @property
    def is_image(self) -> bool:
        return self.image is not None

    def data_url(self) -> str:
        assert self.image is not None
        return "data:image/jpeg;base64," + base64.b64encode(self.image).decode()


def kind_of(filename: str, content: bytes) -> str:
    name = filename.lower()
    ext = name[name.rfind(".") :] if "." in name else ""
    if content[:5] == b"%PDF-" or ext == ".pdf":
        return "pdf"
    if ext in IMAGE_EXT or content[:3] == b"\xff\xd8\xff" or content[:8] == b"\x89PNG\r\n\x1a\n":
        return "image"
    if ext in SHEET_EXT or (content[:2] == b"PK" and ext not in {".docx", ".zip"}):
        return "sheet"
    if ext in TEXT_EXT or ext == "":
        return "text"
    raise AttachmentError(f"{ext or 'this'} files are not supported: use PDF, an image, CSV, TXT or Excel")


def _jpeg(img: "object") -> bytes:
    from PIL import Image

    assert isinstance(img, Image.Image)
    buf = io.BytesIO()
    img.convert("RGB").save(buf, "JPEG", quality=85, optimize=True)
    return buf.getvalue()


def _fit(img: "object", long_side: int = LONG_SIDE) -> "object":
    from PIL import Image

    assert isinstance(img, Image.Image)
    w, h = img.size
    scale = min(1.0, long_side / max(w, h))
    return img.resize((max(1, int(w * scale)), max(1, int(h * scale))), Image.Resampling.LANCZOS) if scale < 1 else img


def image_units(img: "object", label: str) -> list[Unit]:
    """A tall screenshot becomes overlapping tiles so small print stays readable."""
    from PIL import Image

    assert isinstance(img, Image.Image)
    w, h = img.size
    if h <= w * TILE_RATIO * 1.5:
        return [Unit(label, image=_jpeg(_fit(img)))]
    tile = int(w * TILE_RATIO)
    overlap = int(tile * 0.3)  # generous, so a card or table row is whole in at least one tile
    tops = list(range(0, max(1, h - overlap), tile - overlap))
    out = []
    for i, top in enumerate(tops, start=1):
        crop = img.crop((0, top, w, min(h, top + tile)))
        out.append(Unit(f"{label} tile {i}/{len(tops)}", image=_jpeg(_fit(crop))))
    return out


def split(unit: Unit) -> list[Unit]:
    """Halve a unit (by lines, or top/bottom with overlap) when the answer did not fit the output budget."""
    if unit.text is not None:
        lines = unit.text.splitlines()
        if len(lines) < 4:
            return [unit]
        header, body = lines[0], lines[1:]
        mid = len(body) // 2
        return [
            Unit(f"{unit.label} (a)", text="\n".join([header, *body[:mid]])),
            Unit(f"{unit.label} (b)", text="\n".join([header, *body[mid:]])),
        ]
    from PIL import Image

    assert unit.image is not None
    img = Image.open(io.BytesIO(unit.image))
    w, h = img.size
    if h < 300:
        return [unit]
    cut, ov = h // 2, int(h * 0.05)
    return [
        Unit(f"{unit.label} (top)", image=_jpeg(img.crop((0, 0, w, cut + ov)))),
        Unit(f"{unit.label} (bottom)", image=_jpeg(img.crop((0, cut - ov, w, h)))),
    ]


def text_chunks(text: str, label: str) -> list[Unit]:
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if not lines:
        return []
    header, out, cur, size, start = lines[0], [], [lines[0]], len(lines[0]), 1
    for i, ln in enumerate(lines[1:], start=2):
        if size + len(ln) > CHUNK_CHARS and len(cur) > 1:
            out.append(Unit(f"{label} lines {start}-{i - 1}", text="\n".join(cur)))
            cur, size, start = [header], len(header), i
        cur.append(ln)
        size += len(ln) + 1
    out.append(Unit(f"{label} lines {start}-{len(lines)}" if out else label, text="\n".join(cur)))
    return out


def _decode(content: bytes) -> str:
    from app.modules.ingest.adapters.nordnet import decode

    return decode(content)


def load(filename: str, content: bytes) -> list[Unit]:
    if len(content) > MAX_BYTES:
        raise AttachmentError("The file is larger than 25 MB")
    kind = kind_of(filename, content)
    units: list[Unit] = []
    if kind == "pdf":
        import pypdfium2 as pdfium

        try:
            pdf = pdfium.PdfDocument(content)
        except pdfium.PdfiumError as exc:
            raise AttachmentError(f"The PDF could not be opened ({exc}); is it password protected?") from exc
        for i in range(len(pdf)):
            page = pdf[i]
            text = page.get_textpage().get_text_bounded()
            label = f"page {i + 1}"

            def render(page: "pdfium.PdfPage" = page, label: str = label) -> list[Unit]:
                w, h = page.get_size()
                return image_units(page.render(scale=min(LONG_SIDE / max(w, h), 3.0)).to_pil(), label)

            if len(re.sub(r"[^0-9A-Za-zÅÄÖåäö]", "", text)) >= MIN_TEXT:
                parts = text_chunks(text, label) if len(text) > CHUNK_CHARS else [Unit(label, text=text)]
                for part in parts:
                    part.as_image = render
                units.extend(parts)
            else:
                units.extend(render())
    elif kind == "image":
        from PIL import Image, ImageOps

        try:
            img = ImageOps.exif_transpose(Image.open(io.BytesIO(content)))
        except Exception as exc:  # noqa: BLE001 - any decoder error means an unreadable image
            raise AttachmentError(f"The image could not be read ({exc})") from exc
        assert img is not None
        units = image_units(img.convert("RGB"), "image")
    elif kind == "sheet":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise AttachmentError(f"The spreadsheet could not be read ({exc})") from exc
        for ws in wb.worksheets:
            rows = [
                " | ".join("" if c is None else str(c) for c in row)
                for row in ws.iter_rows(values_only=True)
                if any(c not in (None, "") for c in row)
            ]
            units.extend(text_chunks("\n".join(rows), f"sheet {ws.title}"))
    else:
        text = _decode(content)
        first = text.splitlines()[0] if text.strip() else ""
        if "\t" in first:
            reader = csv.reader(io.StringIO(text), delimiter="\t")
            text = "\n".join(" | ".join(r) for r in reader)
        units = text_chunks(text, "text")
    if not units:
        raise AttachmentError("No readable content found in the file")
    if len(units) > MAX_UNITS:
        raise AttachmentError(f"The file has {len(units)} pages or parts; at most {MAX_UNITS} are read at once")
    return units
